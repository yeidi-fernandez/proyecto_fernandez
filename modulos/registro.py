# Registro de eventos de la ejecución.
# Los eventos se guardan en una lista en memoria y al final
# se escriben en un Excel nuevo dentro de logs/.
# No usa la librería logging de Python, como pide el enunciado.

import os
import traceback
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

# Orden de importancia de los niveles. Sirve para filtrar por el nivel mínimo.
NIVELES = {"INFO": 1, "WARNING": 2, "ERROR": 3, "CRITICAL": 4}

COLUMNAS_DETALLE = ["Consecutivo", "Marca de tiempo", "Nivel", "Modulo",
                    "Funcion", "Registro", "Mensaje", "Detalle_tecnico"]

FONDO_ENCABEZADO = PatternFill("solid", fgColor="BDD7EE")
FONDO_ROJO = PatternFill("solid", fgColor="F8CBAD")
FORMATO_HORA = "%Y-%m-%d %H:%M:%S"


def dar_formato_hoja(hoja):
    """Encabezado en negrita con color, fila 1 fija y columnas a la medida."""
    for celda in hoja[1]:
        celda.font = Font(bold=True)
        celda.fill = FONDO_ENCABEZADO
    hoja.freeze_panes = "A2"
    for columna in hoja.columns:
        mas_largo = max(len(str(celda.value or "")) for celda in columna)
        hoja.column_dimensions[columna[0].column_letter].width = min(mas_largo + 2, 80)


class Registro:

    def __init__(self, nivel_minimo="INFO"):
        self.eventos = []
        self.nivel_minimo = nivel_minimo
        self.hora_inicio = datetime.now()

    def registrar(self, nivel, modulo, funcion, mensaje, registro="", detalle=""):
        """Agrega un evento a la lista si su nivel alcanza el mínimo."""
        if NIVELES[nivel] < NIVELES.get(self.nivel_minimo, 1):
            return
        self.eventos.append([
            len(self.eventos) + 1,
            datetime.now().strftime(FORMATO_HORA),
            nivel, modulo, funcion, str(registro), mensaje, detalle,
        ])

    def registrar_error(self, nivel, error, mensaje, registro=""):
        """Registra una excepción. El módulo y la función se sacan del
        recorrido del error: el último punto es donde se lanzó."""
        recorrido = traceback.extract_tb(error.__traceback__)
        if recorrido:
            ultimo = recorrido[-1]
            modulo = os.path.splitext(os.path.basename(ultimo.filename))[0]
            funcion = ultimo.name
        else:
            modulo, funcion = "desconocido", "desconocido"
        detalle = f"{type(error).__name__}: {error}"
        self.registrar(nivel, modulo, funcion, mensaje, registro, detalle)

    def contar_por_nivel(self):
        conteo = {nivel: 0 for nivel in NIVELES}
        for evento in self.eventos:
            conteo[evento[2]] += 1
        return conteo

    def guardar(self, carpeta, prefijo, archivo_entrada="", totales=None):
        """Escribe el Excel del log. Devuelve la ruta del archivo creado."""
        hora_fin = datetime.now()
        totales = totales or {}
        libro = Workbook()

        # Hoja Detalle: una fila por evento
        detalle = libro.active
        detalle.title = "Detalle"
        detalle.append(COLUMNAS_DETALLE)
        for evento in self.eventos:
            detalle.append(evento)
            if evento[2] in ("ERROR", "CRITICAL"):
                for celda in detalle[detalle.max_row]:
                    celda.fill = FONDO_ROJO
        dar_formato_hoja(detalle)

        # Hoja Resumen: datos generales de la ejecución
        resumen = libro.create_sheet("Resumen")
        resumen.append(["Dato", "Valor"])
        resumen.append(["Inicio", self.hora_inicio.strftime(FORMATO_HORA)])
        resumen.append(["Fin", hora_fin.strftime(FORMATO_HORA)])
        resumen.append(["Duración (segundos)",
                        round((hora_fin - self.hora_inicio).total_seconds(), 2)])
        resumen.append(["Archivo de entrada", archivo_entrada or "No se procesó"])
        for nivel, cantidad in self.contar_por_nivel().items():
            resumen.append([f"Eventos {nivel}", cantidad])
        resumen.append(["Novedades recibidas", totales.get("recibidas", 0)])
        resumen.append(["Novedades válidas", totales.get("validas", 0)])
        resumen.append(["Novedades rechazadas", totales.get("rechazadas", 0)])
        dar_formato_hoja(resumen)

        os.makedirs(carpeta, exist_ok=True)
        nombre = f"{prefijo}_{hora_fin.strftime('%Y%m%d_%H%M%S')}.xlsx"
        ruta = os.path.join(carpeta, nombre)
        libro.save(ruta)
        return ruta
    