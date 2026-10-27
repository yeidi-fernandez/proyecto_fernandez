# Carga y valida el Excel de configuración (config/parametros.xlsx).
# Pasa cada hoja a la estructura que más le sirve al programa:
# diccionarios para buscar rápido por código.

import os

from openpyxl import load_workbook

from modulos.excepciones import ErrorConfiguracion

# Cada hoja con las columnas que debe tener.
COLUMNAS_POR_HOJA = {
    "Parametros": ["clave", "valor"],
    "Cargos": ["codigo_cargo", "nombre_cargo", "nivel", "area"],
    "Centros_Costo": ["codigo", "nombre", "area"],
    "Matriz_Roles": ["codigo_cargo", "area", "perfil_sistema", "requiere_aprobacion"],
    "Campos_Obligatorios": ["tipo_novedad", "campos_requeridos"],
}

CLAVES_OBLIGATORIAS = ["dominio_correo", "archivo_salida", "prefijo_log",
                       "formato_fecha_salida", "dias_maximo_futuro",
                       "longitud_telefono", "nivel_log_minimo"]

CLAVES_NUMERICAS = ["dias_maximo_futuro", "longitud_telefono"]


def limpiar(valor):
    """Quita espacios a los textos. Las celdas vacías quedan como texto vacío."""
    if valor is None:
        return ""
    if isinstance(valor, str):
        return valor.strip()
    return valor


def leer_hoja(libro, nombre_hoja):
    """Devuelve las filas de una hoja como lista de diccionarios.
    Las llaves de cada diccionario son los encabezados de la fila 1."""
    filas = list(libro[nombre_hoja].iter_rows(values_only=True))
    if not filas:
        raise ErrorConfiguracion(f"La hoja {nombre_hoja} está vacía")

    encabezados = [str(limpiar(celda)).lower() for celda in filas[0]]
    for columna in COLUMNAS_POR_HOJA[nombre_hoja]:
        if columna not in encabezados:
            raise ErrorConfiguracion(f"A la hoja {nombre_hoja} le falta la columna '{columna}'")

    registros = []
    for fila in filas[1:]:
        if all(limpiar(celda) == "" for celda in fila):
            continue  # fila en blanco, se ignora
        registros.append({encabezados[i]: limpiar(celda) for i, celda in enumerate(fila)})
    return registros


def cargar_parametros(libro):
    parametros = {}
    for fila in leer_hoja(libro, "Parametros"):
        parametros[str(fila["clave"])] = fila["valor"]

    for clave in CLAVES_OBLIGATORIAS:
        if parametros.get(clave, "") == "":
            raise ErrorConfiguracion(f"En la hoja Parametros falta la clave '{clave}'")

    for clave in CLAVES_NUMERICAS:
        try:
            parametros[clave] = int(parametros[clave])
        except (TypeError, ValueError):
            raise ErrorConfiguracion(f"La clave '{clave}' debe ser un número entero")

    # El formato viene escrito como AAAA-MM-DD; Python lo necesita como %Y-%m-%d
    formato = str(parametros["formato_fecha_salida"]).upper()
    for texto, codigo in (("AAAA", "%Y"), ("MM", "%m"), ("DD", "%d")):
        formato = formato.replace(texto, codigo)
    parametros["formato_fecha_python"] = formato

    parametros["nivel_log_minimo"] = str(parametros["nivel_log_minimo"]).upper()
    return parametros


def cargar_configuracion(ruta_archivo):
    """Lee todo el Excel y devuelve un diccionario con cada parte."""
    if not os.path.isfile(ruta_archivo):
        raise ErrorConfiguracion(f"No se encontró el archivo {ruta_archivo}")

    try:
        libro = load_workbook(ruta_archivo, data_only=True)
    except Exception as error:
        raise ErrorConfiguracion(f"No se pudo abrir el archivo de configuración: {error}")

    for hoja in COLUMNAS_POR_HOJA:
        if hoja not in libro.sheetnames:
            raise ErrorConfiguracion(f"Al archivo de configuración le falta la hoja '{hoja}'")

    cargos = {}
    for fila in leer_hoja(libro, "Cargos"):
        cargos[str(fila["codigo_cargo"]).upper()] = {
            "nombre": fila["nombre_cargo"],
            "nivel": fila["nivel"],
            "area": str(fila["area"]).upper(),
        }

    centros = {}
    for fila in leer_hoja(libro, "Centros_Costo"):
        centros[str(fila["codigo"]).upper()] = {
            "nombre": fila["nombre"],
            "area": str(fila["area"]).upper(),
        }

    # La llave es una tupla (cargo, area) para buscar la combinación de una vez
    matriz = {}
    for fila in leer_hoja(libro, "Matriz_Roles"):
        llave = (str(fila["codigo_cargo"]).upper(), str(fila["area"]).upper())
        matriz[llave] = {
            "perfil": fila["perfil_sistema"],
            "requiere_aprobacion": fila["requiere_aprobacion"],
        }

    campos = {}
    for fila in leer_hoja(libro, "Campos_Obligatorios"):
        tipo = str(fila["tipo_novedad"]).upper().replace(" ", "")
        campos[tipo] = [c.strip() for c in str(fila["campos_requeridos"]).split(";") if c.strip()]

    return {
        "parametros": cargar_parametros(libro),
        "cargos": cargos,
        "centros": centros,
        "matriz": matriz,
        "campos": campos,
    }