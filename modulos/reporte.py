# Escribe el Excel de salida en salida/ con sus cuatro hojas.

import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font

from modulos.excepciones import ErrorEscritura
from modulos.registro import dar_formato_hoja

# Orden exacto de columnas que exige el enunciado
COLUMNAS_CARGA = ["consecutivo", "tipo_novedad", "tipo_documento", "numero_documento",
                  "nombre_completo", "correo_corporativo", "codigo_cargo", "nombre_cargo",
                  "nivel", "centro_costo", "area", "perfil_sistema", "requiere_aprobacion",
                  "fecha_efectiva", "telefono", "fecha_vinculacion_previa"]

COLUMNAS_RECHAZOS = ["consecutivo", "tipo", "numero_documento", "nombre",
                     "clasificacion", "campo", "motivo", "valor_recibido"]

COLUMNAS_COLISIONES = ["correo_base", "consecutivo_con_correo_base",
                       "consecutivo_nuevo", "correo_final"]

COLUMNAS_RESUMEN = ["tipo_novedad", "recibidas", "validas", "rechazadas", "porcentaje_exito"]


def llenar_hoja(libro, nombre, columnas, filas):
    """Crea una hoja con su encabezado y una fila por cada diccionario."""
    hoja = libro.create_sheet(nombre)
    hoja.append(columnas)
    for fila in filas:
        hoja.append([fila.get(columna, "") for columna in columnas])
    dar_formato_hoja(hoja)
    return hoja


def llenar_resumen(libro, resumen, archivo_entrada):
    hoja = libro.create_sheet("Resumen")
    hoja.append(COLUMNAS_RESUMEN)
    for tipo, datos in resumen["por_tipo"].items():
        hoja.append([tipo, datos["recibidas"], datos["validas"], datos["rechazadas"], datos["porcentaje"]])

    totales = resumen["totales"]
    hoja.append(["TOTAL", totales["recibidas"], totales["validas"], totales["rechazadas"], totales["porcentaje"]])
    for celda in hoja[hoja.max_row]:
        celda.font = Font(bold=True)

    motivo, cantidad = resumen["motivo_frecuente"]
    hoja.append([])
    hoja.append(["Motivo de rechazo más frecuente", f"{motivo} ({cantidad} veces)"])
    hoja.append(["Archivo de entrada procesado", archivo_entrada])
    for fila in (hoja.max_row - 1, hoja.max_row):
        hoja.cell(row=fila, column=1).font = Font(bold=True)
    dar_formato_hoja(hoja)


def escribir_salida(carpeta, config, validas, rechazados, colisiones, resumen, archivo_entrada):
    """Arma el libro y lo guarda. Devuelve la ruta del archivo creado."""
    libro = Workbook()
    libro.remove(libro.active)  # se quita la hoja vacía que trae por defecto

    llenar_hoja(libro, "Carga_Sistema", COLUMNAS_CARGA, validas)
    llenar_hoja(libro, "Rechazados", COLUMNAS_RECHAZOS, rechazados)
    llenar_hoja(libro, "Colisiones_Correo", COLUMNAS_COLISIONES, colisiones)
    llenar_resumen(libro, resumen, archivo_entrada)

    nombre = f"{config['parametros']['archivo_salida']}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    ruta = os.path.join(carpeta, nombre)
    try:
        libro.save(ruta)
    except PermissionError:
        raise ErrorEscritura(f"No se pudo guardar {nombre}: el archivo está abierto o no hay permiso. "
                             "Ciérrelo y vuelva a ejecutar")
    except OSError as error:
        raise ErrorEscritura(f"No se pudo guardar {nombre}: {error}")
    return ruta